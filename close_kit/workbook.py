"""Configuration workbook (xlsx) -> validated, normalised JSON.

Rules: openpyxl read_only=True; file guards first; any formula cell is an error; every error is
returned with sheet / row / column / reason; if there is one error nothing is accepted.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook

from . import config, safefile
from .audit import FORMULA_START  # = + - @ TAB CR LF and the full-width ones

FIELDS = ("参照", "金額", "勘定日", "摘要")  # fields a pass rule may compare
DISPLAY_FIELDS = ("データ源", "口座", "勘定日", "金額", "参照", "摘要", "元の行番号")
DIRECTIONS = ("入金を+にする", "帳簿の借方を+にする")
FORMATS = ("csv", "zengin", "camt052")
CALENDARS = ("西暦", "令和")
ACCOUNT_CLASSES = ("資産", "負債", "純資産", "収益", "費用")
RULE_KINDS = ("自動", "候補")
RULE_TYPES = ("1対1", "1対多", "多対1", "多対多")
REASON_ACTIONS = ("仕訳", "調整項目")
BANK_SIDES = ("左", "右")
TASK_TYPES = ("繰り返し", "プロジェクト", "臨時")
COUNTINGS = ("営業日", "暦日", "平日")
CONDITIONS = ("なし", "未処理 0 件", "残高一致の差 0")

# (column name, type, required)
Col = tuple[str, str, bool]
SHEETS: dict[str, list[Col]] = {
    "説明": [("schemaの版", "str", True), ("会社名", "str", True), ("作成者", "str", False),
             ("更新日", "date", False), ("会社コード", "str", True), ("仕訳の種類", "str", True),
             ("期日の日数の下限", "int", False), ("期日の日数の上限", "int", False)],
    "会計期間": [("期間の名前", "str", True), ("開始日", "date", True), ("終了日", "date", True),
                 ("締め", "bool", True)],
    "勘定科目": [("勘定コード", "str", True), ("名前", "str", True),
                 ("区分", "enum:" + ",".join(ACCOUNT_CLASSES), True)],
    "データ源": [("名前", "str", True), ("形式", "enum:" + ",".join(FORMATS), True),
                 ("口座", "str", False), ("向き", "enum:" + ",".join(DIRECTIONS), True),
                 ("最小単位の倍率", "posint", True), ("日付の暦", "enum:" + ",".join(CALENDARS), False),
                 ("行のキーの列", "str", False), ("CSV日付列", "str", False),
                 ("CSV金額列", "str", False), ("CSV参照列", "str", False),
                 ("CSV摘要列", "str", False), ("CSV残高列", "str", False)],
    "照合セット": [("名前", "str", True), ("左のデータ源", "str", True),
                   ("右のデータ源", "str", True), ("有効", "bool", True),
                   ("銀行の側", "enum:" + ",".join(BANK_SIDES), True), ("預金の勘定", "str", True),
                   ("残高一致に使う", "bool", True)],
    "表示項目": [("照合セット", "str", True), ("項目", "enum:" + ",".join(DISPLAY_FIELDS), True),
                 ("順番", "posint", True)],
    "小計項目": [("照合セット", "str", True), ("項目", "enum:" + ",".join(DISPLAY_FIELDS), True)],
    "パスルール": [("照合セット", "str", True), ("順番", "posint", True), ("名前", "str", True),
                   ("種別", "enum:" + ",".join(RULE_KINDS), True),
                   ("型", "enum:" + ",".join(RULE_TYPES), True),
                   ("一致させる項目", "str", False), ("date_window", "nonnegint", True),
                   ("amount_tolerance_minor", "nonnegint", True),
                   ("左の束ねるキー", "str", False), ("右の束ねるキー", "str", False),
                   ("束どうしを組にするキー", "str", False), ("1束の行の上限", "posint", False),
                   ("有効", "bool", True), ("差額の相手勘定", "str", False)],
    "理由コード": [("コード", "str", True), ("名前", "str", True),
                   ("処理", "enum:" + ",".join(REASON_ACTIONS), True),
                   ("相手勘定", "str", False), ("有効", "bool", True)],
    "承認ルート": [("下限", "nonnegint", True), ("上限", "posint", False), ("承認できる人", "str", True)],
    "作業一覧": [("作業ID", "str", True), ("名前", "str", True),
                 ("種類", "enum:" + ",".join(TASK_TYPES), True), ("対象の期間", "str", False),
                 ("数え方", "enum:" + ",".join(COUNTINGS), True), ("期末からの日数", "int", True),
                 ("前の作業", "str", False), ("作成者", "str", True), ("確認者", "str", True),
                 ("完了の条件", "enum:" + ",".join(CONDITIONS), True), ("条件の照合セット", "str", False),
                 ("有効", "bool", True)],
    "会社の休日": [("日付", "date", True), ("名前", "str", True)],
}


@dataclass
class WorkbookError:
    sheet: str
    row: int | None
    column: str | None
    message: str

    def __str__(self) -> str:
        where = self.sheet + (f" 行{self.row}" if self.row else "") + (
            f" 列「{self.column}」" if self.column else "")
        return f"{where}: {self.message}"


@dataclass
class WorkbookResult:
    errors: list[WorkbookError] = field(default_factory=list)
    config: dict[str, Any] | None = None
    normalized_json: str = ""
    sha256: str = ""

    @property
    def ok(self) -> bool:
        return not self.errors


def canonical_json(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


class _Reader:
    def __init__(self, errors: list[WorkbookError]):
        self.errors = errors

    def err(self, sheet: str, row: int | None, col: str | None, msg: str) -> None:
        self.errors.append(WorkbookError(sheet, row, col, msg))


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _convert(v: Any, typ: str) -> tuple[bool, Any, str]:
    """Return (ok, value, message)."""
    if typ == "str":
        if isinstance(v, bool):
            return False, None, "文字で入力する"
        if isinstance(v, int):
            return True, str(v), ""
        if isinstance(v, float):
            if v.is_integer():
                return True, str(int(v)), ""
            return True, repr(v), ""
        return True, str(v).strip(), ""
    if typ in ("posint", "nonnegint", "int"):
        if isinstance(v, bool):
            return False, None, "整数で入力する"
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        if isinstance(v, str) and re.fullmatch(r"-?\d+", v.strip()):
            v = int(v.strip())
        if not isinstance(v, int):
            return False, None, f"整数ではない: {v!r}"
        if typ == "posint" and v < 1:
            return False, None, "1 以上の整数にする"
        if typ == "nonnegint" and v < 0:
            return False, None, "0 以上の整数にする"
        return True, v, ""
    if typ == "bool":
        s = str(v).strip()
        if v is True or s in ("はい", "TRUE", "True", "true"):
            return True, True, ""
        if v is False or s in ("いいえ", "FALSE", "False", "false"):
            return True, False, ""
        return False, None, "「はい」か「いいえ」にする"
    if typ == "date":
        if isinstance(v, datetime):
            return True, v.date().isoformat(), ""
        if isinstance(v, date):
            return True, v.isoformat(), ""
        s = str(v).strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
            try:
                return True, datetime.strptime(s, fmt).date().isoformat(), ""
            except ValueError:
                continue
        return False, None, f"日付ではない: {v!r}"
    if typ.startswith("enum:"):
        choices = typ[5:].split(",")
        s = re.sub(r"\s+", "", str(v)).replace("−", "-").replace("＋", "+")
        for c in choices:
            if s == re.sub(r"\s+", "", c):
                return True, c, ""
        return False, None, f"{' / '.join(choices)} のどれかにする（入力: {v!r}）"
    raise AssertionError(typ)


def _peek_schema_version(wb) -> str | None:
    """The value under 「schemaの版」 in the first data row of the 説明 sheet, read before anything else."""
    if "説明" not in wb.sheetnames:
        return None
    rows = list(wb["説明"].iter_rows(values_only=True))
    if len(rows) < 2:
        return None
    head = [str(c).strip().replace(" ", "").replace("　", "") if c is not None else "" for c in rows[0]]
    if "schemaの版" not in head:
        return None
    v = rows[1][head.index("schemaの版")] if head.index("schemaの版") < len(rows[1]) else None
    if v is None:
        return None
    ok, val, _m = _convert(v, "str")
    return val if ok else None


def _split_list(text: str) -> list[str]:
    return [t.strip() for t in re.split(r"[;；,、\n]+", text) if t.strip()]


def _read_sheet(wb, name: str, rd: _Reader) -> list[tuple[int, dict[str, Any]]]:
    """Return [(row number, {column: converted value})]. Reports header, type and formula errors."""
    spec = SHEETS[name]
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=False))
    if not rows:
        rd.err(name, 1, None, "シートが空（見出しの行が無い）")
        return []
    header: dict[str, int] = {}
    for idx, c in enumerate(rows[0]):
        if c.value is not None and getattr(c, "data_type", "") == "f":
            rd.err(name, 1, None, "数式のセルは使えない")
        v = str(c.value).strip().replace(" ", "").replace("　", "") if c.value is not None else ""
        if v:
            header[v] = idx
    for col, _typ, req in spec:
        if col not in header and req:
            rd.err(name, 1, col, "必須の列が無い")
    out: list[tuple[int, dict[str, Any]]] = []
    for ridx, row in enumerate(rows[1:], start=2):
        formula_seen = False
        for c in row:
            if getattr(c, "data_type", "") == "f":
                rd.err(name, ridx, None, "数式のセルは使えない（値だけを入れる）")
                formula_seen = True
                break
        if formula_seen:
            continue
        if all(_is_blank(c.value) for c in row):
            continue
        rec: dict[str, Any] = {}
        for col, typ, req in spec:
            if col not in header:
                rec[col] = None
                continue
            idx = header[col]
            v = row[idx].value if idx < len(row) else None
            if _is_blank(v):
                if req:
                    rd.err(name, ridx, col, "必須の項目が空")
                rec[col] = None
                continue
            ok, val, msg = _convert(v, typ)
            if not ok:
                rd.err(name, ridx, col, msg)
                rec[col] = None
            else:
                rec[col] = val
        out.append((ridx, rec))
    return out


def _parse_keys(text: str | None, sheet: str, row: int, col: str, rd: _Reader) -> list[list[str]]:
    keys: list[list[str]] = []
    if not text:
        return keys
    for tok in _split_list(text):
        if re.search(r"[<>≈~!≠≦≧]", tok):
            rd.err(sheet, row, col, f"比べ方は「=」（等しい）だけ: {tok!r}")
            continue
        m = re.fullmatch(r"(\S+?)\s*=\s*(\S+)", tok)
        if not m:
            rd.err(sheet, row, col, f"「左の項目=右の項目」の形で書く: {tok!r}")
            continue
        l, r = m.group(1), m.group(2)
        for f in (l, r):
            if f not in FIELDS:
                rd.err(sheet, row, col, f"項目は {' / '.join(FIELDS)} のどれか: {f!r}")
        if l in FIELDS and r in FIELDS:
            keys.append([l, r])
    return keys


def _parse_fields(text: str | None, sheet: str, row: int, col: str, rd: _Reader) -> list[str]:
    out: list[str] = []
    for tok in _split_list(text or ""):
        if tok not in FIELDS:
            rd.err(sheet, row, col, f"項目は {' / '.join(FIELDS)} のどれか: {tok!r}")
        else:
            out.append(tok)
    return out


def _find_cycle(graph: dict[str, list[str]]) -> list[str]:
    """A cycle in the predecessor graph as a list of ids (first = last), or []."""
    state: dict[str, int] = {}
    stack: list[str] = []

    def visit(n: str) -> list[str]:
        state[n] = 1
        stack.append(n)
        for m in graph.get(n, []):
            if state.get(m) == 1:
                return stack[stack.index(m):] + [m]
            if state.get(m) is None:
                found = visit(m)
                if found:
                    return found
        stack.pop()
        state[n] = 2
        return []

    for n in graph:
        if state.get(n) is None:
            found = visit(n)
            if found:
                return found
    return []


def _has_balance(src: dict[str, Any]) -> bool:
    return src["format"] in ("zengin", "camt052") or (src["format"] == "csv" and bool(src["csv"]["balance"]))


def _check_reconcile_set(row: int, d: dict[str, Any], sources: list[dict[str, Any]], rd: _Reader) -> None:
    """A match set used for the period-end balance needs both data sources to carry balances, in one unit."""
    if not d["残高一致に使う"] or d["残高一致に使う"] is None:
        return
    by_name = {s["name"]: s for s in sources}
    pair = [by_name.get(d["左のデータ源"]), by_name.get(d["右のデータ源"])]
    if any(p is None for p in pair):
        return  # the missing data source is reported on its own
    for col, src in (("左のデータ源", pair[0]), ("右のデータ源", pair[1])):
        if not _has_balance(src):
            rd.err("照合セット", row, col,
                   f"残高一致に使うには残高が要る（{src['name']}: 形式が全銀・camt052、または CSV に残高の列がある物）")
    if pair[0]["multiplier"] != pair[1]["multiplier"]:
        rd.err("照合セット", row, "残高一致に使う",
               f"2 つのデータ源の最小単位の倍率が違う（{pair[0]['name']}={pair[0]['multiplier']}、"
               f"{pair[1]['name']}={pair[1]['multiplier']}）")


def parse_workbook(data: bytes, filename: str = "workbook.xlsx") -> WorkbookResult:
    """Validate an uploaded workbook. Never raises for content problems; they are returned."""
    errors: list[WorkbookError] = []
    rd = _Reader(errors)
    res = WorkbookResult(errors=errors)
    try:
        safefile.check_xlsx(data, filename)
    except safefile.FileRejected as e:
        errors.append(WorkbookError("(ファイル)", None, None, str(e)))
        return res
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as e:  # noqa: BLE001 - any parser failure is reported as unreadable
        errors.append(WorkbookError("(ファイル)", None, None,
                                    f"xlsx として読めない: {type(e).__name__}"))
        return res
    try:
        old = _peek_schema_version(wb)
        if old is not None and old != config.SCHEMA_VERSION:
            errors.append(WorkbookError(
                "説明", 2, "schemaの版",
                f"schema の版が古い（入力 {old!r}、この版は {config.SCHEMA_VERSION!r}）。"
                "版 2 の形（承認ルート・作業一覧・会社の休日のシートと、説明・照合セットの足した列）で作り直す"))
            return res
        missing = [s for s in SHEETS if s not in wb.sheetnames]
        for s in missing:
            errors.append(WorkbookError(s, None, None, "必要なシートが無い"))
        if missing:
            return res
        data_by_sheet = {s: _read_sheet(wb, s, rd) for s in SHEETS}
    finally:
        wb.close()

    cfg: dict[str, Any] = {}

    # --- 説明
    desc = data_by_sheet["説明"]
    if len(desc) != 1:
        rd.err("説明", None, None, f"データの行は 1 行だけ（今は {len(desc)} 行）")
    elif desc[0][1]["schemaの版"] is not None:
        d = desc[0][1]
        if d["schemaの版"] != config.SCHEMA_VERSION:
            rd.err("説明", desc[0][0], "schemaの版",
                   f"この版では schema の版は {config.SCHEMA_VERSION!r} だけ（入力: {d['schemaの版']!r}）")
        if d["会社コード"] is not None and len(d["会社コード"]) > 4:
            rd.err("説明", desc[0][0], "会社コード", f"4 文字までにする（SAP の BUKRS の桁）: {d['会社コード']!r}")
        if d["仕訳の種類"] is not None and len(d["仕訳の種類"]) > 2:
            rd.err("説明", desc[0][0], "仕訳の種類", f"2 文字までにする（SAP の BLART の桁）: {d['仕訳の種類']!r}")
        for col in ("会社コード", "仕訳の種類"):   # written as they are into the posting file
            if d[col] is not None and d[col].startswith(FORMULA_START):
                rd.err("説明", desc[0][0], col, f"表計算の数式になる文字で始めない: {d[col]!r}")
        lo = d["期日の日数の下限"] if d["期日の日数の下限"] is not None else config.TASK_OFFSET_MIN
        hi = d["期日の日数の上限"] if d["期日の日数の上限"] is not None else config.TASK_OFFSET_MAX
        if lo > 0:
            rd.err("説明", desc[0][0], "期日の日数の下限", f"0 以下にする（0 日目は必ず含む）: {lo}")
        if hi < 0:
            rd.err("説明", desc[0][0], "期日の日数の上限", f"0 以上にする（0 日目は必ず含む）: {hi}")
        cfg["description"] = {"schema_version": d["schemaの版"], "company": d["会社名"],
                              "author": d["作成者"], "updated": d["更新日"],
                              "company_code": d["会社コード"], "journal_type": d["仕訳の種類"],
                              "task_offset_min": lo, "task_offset_max": hi}

    # --- 会計期間
    periods = []
    seen: set[str] = set()
    for r, d in data_by_sheet["会計期間"]:
        if d["期間の名前"] is None or d["開始日"] is None or d["終了日"] is None:
            continue
        if d["期間の名前"] in seen:
            rd.err("会計期間", r, "期間の名前", f"名前が重複: {d['期間の名前']}")
        seen.add(d["期間の名前"])
        if d["終了日"] < d["開始日"]:
            rd.err("会計期間", r, "終了日", "終了日が開始日より前")
        periods.append({"name": d["期間の名前"], "start": d["開始日"], "end": d["終了日"],
                        "closed": bool(d["締め"])})
    periods.sort(key=lambda p: (p["start"], p["name"]))
    for a, b in zip(periods, periods[1:]):
        if b["start"] <= a["end"]:
            rd.err("会計期間", None, None, f"期間が重なっている: {a['name']} と {b['name']}")
    cfg["periods"] = periods

    # --- 勘定科目
    accounts = []
    codes: set[str] = set()
    for r, d in data_by_sheet["勘定科目"]:
        if d["勘定コード"] is None:
            continue
        if len(d["勘定コード"]) > 10:
            rd.err("勘定科目", r, "勘定コード", f"10 文字までにする（SAP の HKONT の桁）: {d['勘定コード']!r}")
        if d["勘定コード"].startswith(FORMULA_START):   # written as it is into the posting file
            rd.err("勘定科目", r, "勘定コード", f"表計算の数式になる文字で始めない: {d['勘定コード']!r}")
        if d["勘定コード"] in codes:
            rd.err("勘定科目", r, "勘定コード", f"コードが重複: {d['勘定コード']}")
        codes.add(d["勘定コード"])
        accounts.append({"code": d["勘定コード"], "name": d["名前"], "class": d["区分"]})
    accounts.sort(key=lambda a: a["code"])
    cfg["accounts"] = accounts
    account_class = {a["code"]: a["class"] for a in accounts}

    # --- データ源
    sources = []
    names: set[str] = set()
    for r, d in data_by_sheet["データ源"]:
        if d["名前"] is None or d["形式"] is None:
            continue
        if d["名前"] in names:
            rd.err("データ源", r, "名前", f"名前が重複: {d['名前']}")
        names.add(d["名前"])
        fmt = d["形式"]
        if fmt in ("zengin", "camt052"):
            if not d["口座"]:
                rd.err("データ源", r, "口座", f"形式 {fmt} では口座が必要")
            if fmt == "zengin" and not d["日付の暦"]:
                rd.err("データ源", r, "日付の暦", "全銀では日付の暦（西暦・令和）が必要")
        if fmt in ("zengin", "camt052") and d["向き"] != "入金を+にする":
            rd.err("データ源", r, "向き",
                   "全銀・camt052 は入金/出金の区分で符号が決まるので「入金を+にする」にする")
        if fmt == "csv":
            for col in ("CSV日付列", "CSV金額列"):
                if not d[col]:
                    rd.err("データ源", r, col, "CSV では日付と金額の列名が必要")
        sources.append({
            "name": d["名前"], "format": fmt, "account": d["口座"] or "",
            "direction": d["向き"], "multiplier": d["最小単位の倍率"],
            "calendar": d["日付の暦"] or "西暦",
            "row_key_cols": _split_list(d["行のキーの列"] or ""),
            "csv": {"date": d["CSV日付列"] or "", "amount": d["CSV金額列"] or "",
                    "reference": d["CSV参照列"] or "", "memo": d["CSV摘要列"] or "",
                    "balance": d["CSV残高列"] or ""},
        })
    sources.sort(key=lambda s: s["name"])
    cfg["sources"] = sources

    # --- 照合セット
    match_sets = []
    ms_names: set[str] = set()
    for r, d in data_by_sheet["照合セット"]:
        if d["名前"] is None:
            continue
        if d["名前"] in ms_names:
            rd.err("照合セット", r, "名前", f"名前が重複: {d['名前']}")
        ms_names.add(d["名前"])
        for col in ("左のデータ源", "右のデータ源"):
            if d[col] is not None and d[col] not in names:
                rd.err("照合セット", r, col, f"データ源に無い名前: {d[col]}")
        if d["左のデータ源"] and d["左のデータ源"] == d["右のデータ源"]:
            rd.err("照合セット", r, "右のデータ源", "左と右が同じデータ源（2 者の照合だけ作る）")
        dep = d["預金の勘定"]
        if dep is not None and dep not in codes:
            rd.err("照合セット", r, "預金の勘定", f"勘定科目に無いコード: {dep}")
        elif dep is not None and account_class.get(dep) != "資産":
            rd.err("照合セット", r, "預金の勘定", f"区分が資産の勘定にする（{dep} は {account_class.get(dep)}）")
        match_sets.append({"name": d["名前"], "left": d["左のデータ源"], "right": d["右のデータ源"],
                           "enabled": bool(d["有効"]), "bank_side": d["銀行の側"],
                           "deposit_account": dep or "", "reconcile": bool(d["残高一致に使う"])})
        _check_reconcile_set(r, d, sources, rd)
    match_sets.sort(key=lambda m: m["name"])
    seen_dep: dict[str, str] = {}
    for m in match_sets:
        if m["reconcile"] and m["deposit_account"]:
            if m["deposit_account"] in seen_dep:
                rd.err("照合セット", None, "預金の勘定",
                       f"残高一致に使う照合セットの間で預金の勘定が重なる（{seen_dep[m['deposit_account']]} と "
                       f"{m['name']} が {m['deposit_account']}）。重なると仕訳を 2 回足す")
            seen_dep[m["deposit_account"]] = m["name"]
    cfg["match_sets"] = match_sets

    # --- 表示項目・小計項目
    display = []
    for r, d in data_by_sheet["表示項目"]:
        if d["照合セット"] is None:
            continue
        if d["照合セット"] not in ms_names:
            rd.err("表示項目", r, "照合セット", f"照合セットに無い名前: {d['照合セット']}")
        display.append({"match_set": d["照合セット"], "field": d["項目"], "order": d["順番"]})
    display.sort(key=lambda x: (x["match_set"], x["order"] or 0, x["field"] or ""))
    cfg["display_fields"] = display
    subtotals = []
    for r, d in data_by_sheet["小計項目"]:
        if d["照合セット"] is None:
            continue
        if d["照合セット"] not in ms_names:
            rd.err("小計項目", r, "照合セット", f"照合セットに無い名前: {d['照合セット']}")
        subtotals.append({"match_set": d["照合セット"], "field": d["項目"]})
    subtotals.sort(key=lambda x: (x["match_set"], x["field"] or ""))
    cfg["subtotal_fields"] = subtotals

    # --- パスルール
    rules = []
    seen_orders: set[tuple[str, int]] = set()
    for r, d in data_by_sheet["パスルール"]:
        if d["照合セット"] is None or d["順番"] is None:
            continue
        if d["照合セット"] not in ms_names:
            rd.err("パスルール", r, "照合セット", f"照合セットに無い名前: {d['照合セット']}")
        key = (d["照合セット"], d["順番"])
        if key in seen_orders:
            rd.err("パスルール", r, "順番", f"順番が重なっている（照合セット {d['照合セット']} の {d['順番']}）")
        seen_orders.add(key)
        typ = d["型"]
        keys = _parse_keys(d["一致させる項目"], "パスルール", r, "一致させる項目", rd)
        lb = _parse_fields(d["左の束ねるキー"], "パスルール", r, "左の束ねるキー", rd)
        rb = _parse_fields(d["右の束ねるキー"], "パスルール", r, "右の束ねるキー", rd)
        pk = d["束どうしを組にするキー"]
        if pk is not None and pk not in FIELDS:
            rd.err("パスルール", r, "束どうしを組にするキー", f"項目は {' / '.join(FIELDS)} のどれか: {pk!r}")
            pk = None
        limit = d["1束の行の上限"]
        diff_acc = d["差額の相手勘定"] or ""
        if diff_acc:
            if d["種別"] != "自動" or typ != "1対1":
                rd.err("パスルール", r, "差額の相手勘定",
                       "差額の相手勘定は種別「自動」・型「1対1」の規則だけに書く（許容差を使えるのは 1対1 だけ）")
            if diff_acc not in codes:
                rd.err("パスルール", r, "差額の相手勘定", f"勘定科目に無いコード: {diff_acc}")
            elif diff_acc in {m["deposit_account"] for m in match_sets if m["reconcile"]} or diff_acc in {
                    m["deposit_account"] for m in match_sets if m["name"] == d["照合セット"]}:
                rd.err("パスルール", r, "差額の相手勘定",
                       f"預金の勘定（{diff_acc}）は相手勘定にできない（預金と預金の仕訳になり、差を直さない）")
            elif diff_acc.startswith(FORMULA_START):
                rd.err("パスルール", r, "差額の相手勘定", f"表計算の数式になる文字で始めない: {diff_acc!r}")
        if typ == "1対1":
            if lb or rb or pk or limit:
                rd.err("パスルール", r, None, "1対1 では束ねるキー・組にするキー・上限は使わない")
        else:
            # the type must agree with which bundle key is filled (1対多: right only; 多対1: left only)
            if keys:
                rd.err("パスルール", r, "一致させる項目",
                       f"{typ} では一致させる項目は使わない（束どうしを組にするキーで組にする）")
            if not pk:
                rd.err("パスルール", r, "束どうしを組にするキー", f"{typ} では組にするキーが必要")
            if not limit:
                rd.err("パスルール", r, "1束の行の上限", f"{typ} では 1 束の行の上限が必要")
            want_left = typ in ("多対1", "多対多")
            want_right = typ in ("1対多", "多対多")
            if want_left and not lb:
                rd.err("パスルール", r, "左の束ねるキー", f"{typ} では左の束ねるキーが必要")
            if not want_left and lb:
                rd.err("パスルール", r, "左の束ねるキー",
                       f"{typ} では左は 1 行ずつ。左の束ねるキーは空にする（型と束ねるキーが合わない）")
            if want_right and not rb:
                rd.err("パスルール", r, "右の束ねるキー", f"{typ} では右の束ねるキーが必要")
            if not want_right and rb:
                rd.err("パスルール", r, "右の束ねるキー",
                       f"{typ} では右は 1 行ずつ。右の束ねるキーは空にする（型と束ねるキーが合わない）")
        rules.append({
            "match_set": d["照合セット"], "order": d["順番"], "name": d["名前"], "kind": d["種別"],
            "type": typ, "keys": keys, "date_window": d["date_window"],
            "amount_tolerance_minor": d["amount_tolerance_minor"],
            "left_bundle": lb, "right_bundle": rb, "pair_key": pk, "bundle_limit": limit,
            "enabled": bool(d["有効"]), "diff_account": diff_acc,
        })
    rules.sort(key=lambda x: (x["match_set"], x["order"]))
    cfg["pass_rules"] = rules

    # --- 理由コード
    reasons = []
    rc: set[str] = set()
    for r, d in data_by_sheet["理由コード"]:
        if d["コード"] is None:
            continue
        if d["コード"] in rc:
            rd.err("理由コード", r, "コード", f"コードが重複: {d['コード']}")
        rc.add(d["コード"])
        if d["処理"] == "仕訳" and not d["相手勘定"]:
            rd.err("理由コード", r, "相手勘定", "処理が「仕訳」のときは相手勘定が必要")
        if d["相手勘定"] and d["相手勘定"] not in codes:
            rd.err("理由コード", r, "相手勘定", f"勘定科目に無いコード: {d['相手勘定']}")
        elif d["相手勘定"] and d["相手勘定"] in {m["deposit_account"] for m in match_sets if m["reconcile"]}:
            rd.err("理由コード", r, "相手勘定", f"残高一致に使う預金の勘定（{d['相手勘定']}）は相手勘定にできない"
                   "（預金と預金の仕訳になり、差を直さない）")
        reasons.append({"code": d["コード"], "name": d["名前"], "action": d["処理"],
                        "counter_account": d["相手勘定"] or "", "enabled": bool(d["有効"])})
    reasons.sort(key=lambda x: x["code"])
    cfg["reason_codes"] = reasons

    # --- 承認ルート: amount ranges from 0 upward, no gap, no overlap, only the last upper limit empty
    routes = []
    for r, d in data_by_sheet["承認ルート"]:
        if d["下限"] is None and d["承認できる人"] is None:
            continue
        names_ = _split_list(d["承認できる人"] or "")
        if d["承認できる人"] is not None and not names_:
            rd.err("承認ルート", r, "承認できる人", "名前が 1 人も無い")
        if d["下限"] is not None:
            routes.append({"lower": d["下限"], "upper": d["上限"], "approvers": names_, "row": r})
    routes.sort(key=lambda x: (x["lower"], x["upper"] if x["upper"] is not None else 1 << 62))
    if not routes and not any(e.sheet == "承認ルート" for e in errors):
        rd.err("承認ルート", None, None, "行が 1 つも無い（承認できる人を決める幅が要る）")
    for i, x in enumerate(routes):
        last = i == len(routes) - 1
        if i == 0 and x["lower"] != 0:
            rd.err("承認ルート", x["row"], "下限", f"最初の幅の下限は 0 にする（今は {x['lower']}）")
        if i > 0:
            prev = routes[i - 1]
            if prev["upper"] is not None and x["lower"] > prev["upper"]:
                rd.err("承認ルート", x["row"], "下限",
                       f"幅に隙間がある（前の幅の上限 {prev['upper']} と下限 {x['lower']}）")
            elif prev["upper"] is not None and x["lower"] < prev["upper"]:
                rd.err("承認ルート", x["row"], "下限",
                       f"幅が重なっている（前の幅の上限 {prev['upper']} と下限 {x['lower']}）")
        if not last and x["upper"] is None:
            rd.err("承認ルート", x["row"], "上限", "上限が空なのは最後の行だけ")
        if last and x["upper"] is not None:
            rd.err("承認ルート", x["row"], "上限", "最後の行の上限は空にする（上限なし）")
        if x["upper"] is not None and x["upper"] <= x["lower"]:
            rd.err("承認ルート", x["row"], "上限", f"上限が下限以下（下限 {x['lower']}、上限 {x['upper']}）")
    cfg["approval_routes"] = [{"lower": x["lower"], "upper": x["upper"], "approvers": x["approvers"]}
                              for x in routes]

    # --- 会社の休日
    holidays = []
    hseen: set[str] = set()
    for r, d in data_by_sheet["会社の休日"]:
        if d["日付"] is None:
            continue
        if d["日付"] in hseen:
            rd.err("会社の休日", r, "日付", f"日付が重複: {d['日付']}")
        hseen.add(d["日付"])
        holidays.append({"date": d["日付"], "name": d["名前"] or ""})
    holidays.sort(key=lambda x: x["date"])
    cfg["company_holidays"] = holidays

    # --- 作業一覧
    tasks = []
    task_ids: set[str] = set()
    period_names = {p["name"] for p in periods}
    ms_by_name = {m["name"]: m for m in match_sets}
    rows_by_id: dict[str, int] = {}
    for r, d in data_by_sheet["作業一覧"]:
        if d["作業ID"] is None:
            continue
        if d["作業ID"] in task_ids:
            rd.err("作業一覧", r, "作業ID", f"作業 ID が重複: {d['作業ID']}")
        task_ids.add(d["作業ID"])
        rows_by_id[d["作業ID"]] = r
        typ = d["種類"]
        if typ == "繰り返し" and d["対象の期間"]:
            rd.err("作業一覧", r, "対象の期間", "繰り返しの作業では対象の期間は使わない（全部の締めていない期間に作る）")
        if typ in ("プロジェクト", "臨時"):
            if not d["対象の期間"]:
                rd.err("作業一覧", r, "対象の期間", f"{typ}の作業では対象の期間が必要")
            elif d["対象の期間"] not in period_names:
                rd.err("作業一覧", r, "対象の期間", f"会計期間に無い名前: {d['対象の期間']}")
        if d["期末からの日数"] is not None and "description" in cfg:
            lo, hi = cfg["description"]["task_offset_min"], cfg["description"]["task_offset_max"]
            if not lo <= d["期末からの日数"] <= hi:
                rd.err("作業一覧", r, "期末からの日数", f"{lo} から {hi} の整数にする（範囲は説明シートの設定）")
        if d["作成者"] and d["確認者"] and d["作成者"] == d["確認者"]:
            rd.err("作業一覧", r, "確認者", f"作成者と確認者が同じ人（{d['作成者']}）。別の人にする")
        cond = d["完了の条件"]
        if cond and cond != "なし":
            if not d["条件の照合セット"]:
                rd.err("作業一覧", r, "条件の照合セット", f"完了の条件「{cond}」では照合セットが必要")
            elif d["条件の照合セット"] not in ms_by_name:
                rd.err("作業一覧", r, "条件の照合セット", f"照合セットに無い名前: {d['条件の照合セット']}")
            elif cond == "残高一致の差 0" and not ms_by_name[d["条件の照合セット"]]["reconcile"]:
                rd.err("作業一覧", r, "条件の照合セット",
                       f"{d['条件の照合セット']} は「残高一致に使う」が「はい」ではない")
        tasks.append({"id": d["作業ID"], "name": d["名前"], "type": typ, "period": d["対象の期間"] or "",
                      "counting": d["数え方"], "offset": d["期末からの日数"],
                      "predecessors": _split_list(d["前の作業"] or ""), "preparer": d["作成者"],
                      "reviewer": d["確認者"], "condition": cond,
                      "condition_match_set": d["条件の照合セット"] or "", "enabled": bool(d["有効"])})
    graph: dict[str, list[str]] = {}
    for t in tasks:
        for pre in t["predecessors"]:
            if pre not in task_ids:
                rd.err("作業一覧", rows_by_id[t["id"]], "前の作業", f"作業 ID に無い名前: {pre}")
            elif pre == t["id"]:
                rd.err("作業一覧", rows_by_id[t["id"]], "前の作業", "自分自身は前の作業にできない")
        graph[t["id"]] = [p for p in t["predecessors"] if p in task_ids]
    cyc = _find_cycle(graph)
    if cyc:
        rd.err("作業一覧", rows_by_id.get(cyc[0]), "前の作業", "前後の順番が輪になっている: " + " → ".join(cyc))
    tasks.sort(key=lambda t: t["id"])
    cfg["tasks"] = tasks

    if errors:
        return res
    res.config = cfg
    res.normalized_json = canonical_json(cfg)
    res.sha256 = hashlib.sha256(res.normalized_json.encode("utf-8")).hexdigest()
    return res
