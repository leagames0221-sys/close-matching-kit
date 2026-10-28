"""Change requests: moving a version of the settings into production.

A version is made by an administrator from a workbook in the sandbox. Production is pointed at a version only by an
applied change request. The flow (NIST SP 800-128, "Configuration change control": record, decide whether it needs
control, analyse the impact, test, approve, implement, verify, close out; the form is appendix E):

    下書き -> 事前承認待ち -> (the test runs by itself) 試験済み -> 結果の承認済み -> 本番に適用済み -> 確認済み
    (却下 / 取り下げ end it before it is applied; 確認できない keeps it open and lets a rollback request be drafted)

People (checked here, on the server, whatever the screen shows - ADR 0008):
 - the requester can never pre-approve, reject, approve a result or apply, whatever roles he has;
 - the result needs two people who are not the requester: one with the role 承認者 and one with the role 管理者
   (one person with both roles counts once); an approval belongs to the test result it was given for, so a new test
   voids the earlier approvals;
 - the requester confirms that the change came out as asked.

The row of a request is frozen when it leaves 下書き (a trigger); every step is a row of `change_request_events`, the test
results are rows of `change_tests`, and the approvals of a result are rows of `change_approvals` (all append-only).
"""
from __future__ import annotations

import csv
import functools
import io
import json
import sqlite3
from datetime import date
from typing import Any, Callable

from openpyxl import load_workbook

from . import audit, changecheck, changeinfo, config, roles, versions, workbook
from .roles import User
from .timeutil import now_iso

PROD = config.ENV_PRODUCTION


class ChangeError(Exception):
    """The request cannot be carried out (the message is shown to the person)."""


class NotAllowedToChange(ChangeError):
    """The right role, the wrong person (the requester decides his own request, a stranger confirms)."""


class StaleTest(ChangeError):
    """Production moved after the test: the request goes back to 試験待ち."""


class NotLatestResult(ChangeError):
    """The person decides on a test result that is no longer the latest one (the test ran again after he opened the
    page): his approval would count for a result he has not seen."""


# ---- small helpers ---------------------------------------------------------------------------------------------------

def _req(conn: sqlite3.Connection, rid: int) -> sqlite3.Row:
    r = conn.execute("SELECT * FROM change_requests WHERE id = ?", (rid,)).fetchone()
    if r is None:
        raise ChangeError(f"変更の申請 {rid} は無い")
    return r


def state_of(conn: sqlite3.Connection, rid: int) -> str:
    r = conn.execute("SELECT to_state FROM change_request_events WHERE request_id = ? ORDER BY id DESC LIMIT 1",
                     (rid,)).fetchone()
    if r is None:
        raise ChangeError(f"変更の申請 {rid} の履歴が無い")
    return r["to_state"]


def _event(conn: sqlite3.Connection, user: User, rid: int, action: str, state: str, reason: str = "",
           version_id: int | None = None, env: str = PROD) -> None:
    """Inside the caller's transaction."""
    conn.execute(
        "INSERT INTO change_request_events(request_id, at, user_id, user_name, roles, env, version_id, action, to_state,"
        " reason) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (rid, now_iso(), user.id, user.name, ",".join(sorted(user.roles)), env, version_id, action, state,
         reason[:1000]))


def _denied(conn: sqlite3.Connection, user: User, action: str, rid: int | None, detail: str) -> None:
    audit.log(conn, user_id=user.id, user_name=user.name, action=action, outcome=audit.OUTCOME_DENIED, env=PROD,
              detail=(f"申請 {rid}: " if rid else "") + detail)


def _ok(conn: sqlite3.Connection, user: User, action: str, rid: int, detail: str, version_id: int | None = None
        ) -> None:
    audit.log(conn, user_id=user.id, user_name=user.name, action=action, outcome=audit.OUTCOME_OK, env=PROD,
              version_id=version_id, detail=f"申請 {rid}: {detail}")


def _refusals_logged(action: str, rid_first: bool = True) -> Callable:
    """Every refusal of the step becomes a row of the operation log (outcome 拒否, with the reason) - also the refusals of
    the state ("「却下」で、適用できる状態ではない"), of an empty field and of a second approval, which the step itself does not
    write. A refusal that the step has written already (the requester, a closed period, the sha256) is not written twice."""
    def deco(fn: Callable) -> Callable:
        @functools.wraps(fn)
        def wrapper(conn: sqlite3.Connection, user: User, *args: Any, **kwargs: Any) -> Any:
            mark = conn.execute("SELECT COALESCE(MAX(id), 0) FROM audit_log").fetchone()[0]
            try:
                return fn(conn, user, *args, **kwargs)
            except ChangeError as e:
                written = conn.execute("SELECT 1 FROM audit_log WHERE id > ? AND user_id IS ? AND outcome = ?",
                                       (mark, user.id, audit.OUTCOME_DENIED)).fetchone()
                if written is None:
                    rid = args[0] if rid_first and args and isinstance(args[0], int) else None
                    _denied(conn, user, action, rid, str(e))
                raise
        return wrapper
    return deco


def _not_requester(conn: sqlite3.Connection, user: User, req: sqlite3.Row, action: str) -> None:
    """Whatever roles the user holds: the requester never decides his own request."""
    if req["requester_id"] == user.id:
        msg = "申請者は、自分の申請の承認・却下・適用ができない（役割をいくつ持っていても）"
        _denied(conn, user, action, req["id"], msg)
        raise NotAllowedToChange(msg)


def _need_state(conn: sqlite3.Connection, req: sqlite3.Row, allowed: tuple[str, ...], what: str) -> str:
    st = state_of(conn, req["id"])
    if st not in allowed:
        raise ChangeError(f"申請 {req['id']} は「{st}」で、{what}できる状態ではない")
    return st


def _version(conn: sqlite3.Connection, vid: int) -> tuple[sqlite3.Row, dict[str, Any]]:
    row = conn.execute("SELECT * FROM config_versions WHERE id = ?", (vid,)).fetchone()
    if row is None:
        raise ChangeError(f"設定の版 {vid} は無い")
    return row, json.loads(row["normalized_json"])


def _ver(vid: int | None) -> str:
    """How a settings version reads in a sentence: 版 3, or 版なし when there is none (as on the screens)."""
    return "版なし" if vid is None else f"版 {vid}"


def production_version(conn: sqlite3.Connection) -> int | None:
    return versions.env_version_id(conn, PROD)


def check_not_reopening(conn: sqlite3.Connection, user: User, to_cfg: dict[str, Any], action: str,
                        rid: int | None = None) -> None:
    """A version that opens a period that production has closed is refused. The base is the version in PRODUCTION (the
    check at the loading of a workbook compares with the sandbox)."""
    pv = production_version(conn)
    if pv is None:
        return
    _row, base = _version(conn, pv)
    reopened = versions.reopened_periods(base, to_cfg)
    if reopened:
        msg = "本番で締めた期間を開く版は申請も適用もできない: " + "、".join(reopened)
        _denied(conn, user, action, rid, msg)
        raise ChangeError(msg)


def _check_target(conn: sqlite3.Connection, user: User, to_version_id: int, action: str) -> dict[str, Any]:
    _row, cfg = _version(conn, to_version_id)
    if cfg.get("description", {}).get("schema_version") != config.SCHEMA_VERSION:
        raise ChangeError("schema の版が古い版は本番に上げられない（戻し先にもできない）")
    if to_version_id == production_version(conn):
        raise ChangeError(f"設定の版 {to_version_id} はすでに本番の版")
    check_not_reopening(conn, user, cfg, action)
    return cfg


# ---- making a request ------------------------------------------------------------------------------------------------

@_refusals_logged("change_create", rid_first=False)
def create_draft(conn: sqlite3.Connection, user: User, *, title: str, to_version_id: int, justification: str,
                 functional_impact: str = "", impact_if_not: str = "", rollback_version_id: int | None = None,
                 kind: str = "通常", rollback_of: int | None = None) -> int:
    roles.require(conn, user, "change_create", env=PROD)
    if not title.strip():
        raise ChangeError("題名が空")
    if not justification.strip():
        raise ChangeError("理由が空")
    _check_target(conn, user, to_version_id, "change_create")
    if rollback_version_id is not None:
        _row, rcfg = _version(conn, rollback_version_id)
        if rcfg.get("description", {}).get("schema_version") != config.SCHEMA_VERSION:
            raise ChangeError("schema の版が古い版は戻し先にできない")
    with conn:
        cur = conn.execute(
            "INSERT INTO change_requests(kind, status, title, requester_id, created_at, from_version_id,"
            " to_version_id, justification, urgency, functional_impact, impact_if_not, rollback_version_id,"
            " rollback_of) VALUES (?, '下書き', ?, ?, ?, ?, ?, ?, 'Scheduled', ?, ?, ?, ?)",
            (kind, title.strip()[:200], user.id, now_iso(), production_version(conn), to_version_id,
             justification.strip()[:2000], functional_impact.strip()[:2000], impact_if_not.strip()[:2000],
             rollback_version_id, rollback_of))
        rid = int(cur.lastrowid)
        _event(conn, user, rid, "作成", config.CR_DRAFT, version_id=to_version_id)
    _ok(conn, user, "change_create", rid, f"下書きを作った（設定の版 {to_version_id} を本番へ）", to_version_id)
    return rid


def _closed_merged_workbook(source: bytes, prod_cfg: dict[str, Any]) -> bytes:
    """The workbook of the rollback target with the periods that production has closed marked 締め = はい (periods the
    target does not know are added). A real workbook, so the version made from it reads back exactly as it is stored."""
    wb = load_workbook(io.BytesIO(source))
    try:
        ws = wb["会計期間"]
        head = {str(c.value).strip().replace(" ", "").replace("　", ""): i for i, c in enumerate(ws[1]) if c.value}
        need = ("期間の名前", "開始日", "終了日", "締め")
        if any(k not in head for k in need):
            raise ChangeError("戻し先の版の「会計期間」シートの見出しが読めない")
        closed = {p["name"]: p for p in prod_cfg.get("periods", []) if p["closed"]}
        seen: set[str] = set()
        for row in ws.iter_rows(min_row=2):
            name_cell = row[head["期間の名前"]]
            if name_cell.value is None:
                continue
            name = str(name_cell.value).strip()
            seen.add(name)
            if name in closed:
                row[head["締め"]].value = "はい"
        for name, p in closed.items():
            if name not in seen:
                vals: list[Any] = [None] * (max(head.values()) + 1)
                vals[head["期間の名前"]] = name
                vals[head["開始日"]] = date.fromisoformat(p["start"])
                vals[head["終了日"]] = date.fromisoformat(p["end"])
                vals[head["締め"]] = "はい"
                ws.append(vals)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    finally:
        wb.close()


@_refusals_logged("change_create", rid_first=False)
def create_rollback(conn: sqlite3.Connection, user: User, target_version_id: int, *, title: str = "",
                    justification: str, source_request_id: int | None = None) -> int:
    """Draft a request that puts `target_version_id` back - with the periods that production has closed now added to
    it as closed (putting the old version back as it was would open them). The new version is stored; the draft goes
    through the same steps as any request."""
    roles.require(conn, user, "change_create", env=PROD)
    pv = production_version(conn)
    if pv is None:
        raise ChangeError("本番に版が無いので、戻す物が無い")
    trow, tcfg = _version(conn, target_version_id)
    if tcfg.get("description", {}).get("schema_version") != config.SCHEMA_VERSION:
        raise ChangeError("schema の版が古い版は戻し先にできない")
    if target_version_id == pv:
        raise ChangeError(f"設定の版 {target_version_id} はすでに本番の版")
    _prow, pcfg = _version(conn, pv)
    data = _closed_merged_workbook(bytes(trow["source_bytes"]), pcfg)
    res = workbook.parse_workbook(data, "rollback.xlsx")
    if not res.ok:
        raise ChangeError("戻し先の版に今の締めを足した版が作れない: " + "; ".join(str(e) for e in res.errors[:3]))
    with conn:
        new_vid = versions.add_version(conn, user, res, data)
    audit.log(conn, user_id=user.id, user_name=user.name, action="change_rollback_version",
              outcome=audit.OUTCOME_OK, env=PROD, version_id=new_vid,
              detail=f"設定の版 {target_version_id} に、本番の今の締めを足した版 {new_vid} を作った")
    return create_draft(conn, user, title=title.strip() or f"戻し: 設定の版 {target_version_id} へ",
                        to_version_id=new_vid, justification=justification, kind="戻し",
                        rollback_version_id=pv, rollback_of=source_request_id,
                        functional_impact="本番の設定を、版 %d の中身（本番の今の締めを足した物）に戻す" % target_version_id)


@_refusals_logged("change_submit")
def submit(conn: sqlite3.Connection, user: User, rid: int) -> str:
    """Put a draft forward: the parts the form fills by itself (what changes, what it touches, the control impact) are
    computed now and the row is frozen. Only the requester."""
    roles.require(conn, user, "change_create", env=PROD)
    req = _req(conn, rid)
    _need_state(conn, req, (config.CR_DRAFT,), "申請")
    if req["requester_id"] != user.id:
        _denied(conn, user, "change_submit", rid, "作った人だけが申請できる")
        raise NotAllowedToChange("申請できるのは、その下書きを作った人だけ")
    to_cfg = _check_target(conn, user, req["to_version_id"], "change_submit")
    pv = production_version(conn)
    base_cfg = _version(conn, pv)[1] if pv is not None else None
    diff = changeinfo.diff_configs(base_cfg, to_cfg)
    imp = changeinfo.impact(conn, base_cfg, to_cfg, PROD)
    sec = changeinfo.security_impact(diff, base_cfg, to_cfg)
    with conn:
        conn.execute(
            "UPDATE change_requests SET status = '申請済み', from_version_id = ?, submitted_at = ?, diff_json = ?,"
            " impact_json = ?, security_json = ?, rollback_version_id = COALESCE(rollback_version_id, ?) WHERE id = ?",
            (pv, now_iso(), json.dumps(diff, ensure_ascii=False), json.dumps(imp, ensure_ascii=False),
             json.dumps(sec, ensure_ascii=False), pv, rid))
        _event(conn, user, rid, "申請", config.CR_WAIT_PRE, version_id=req["to_version_id"])
    _ok(conn, user, "change_submit", rid, "申請した（事前承認待ち）", req["to_version_id"])
    return config.CR_WAIT_PRE


@_refusals_logged("change_withdraw")
def withdraw(conn: sqlite3.Connection, user: User, rid: int, reason: str = "") -> str:
    roles.require(conn, user, "change_withdraw", env=PROD)
    req = _req(conn, rid)
    if req["requester_id"] != user.id:
        _denied(conn, user, "change_withdraw", rid, "作った人だけが取り下げられる")
        raise NotAllowedToChange("取り下げできるのは、その申請を作った人だけ")
    _need_state(conn, req, (config.CR_DRAFT, config.CR_WAIT_PRE, config.CR_WAIT_TEST, config.CR_TESTED,
                            config.CR_RESULT_OK), "取り下げ")
    with conn:
        if req["status"] == "下書き":      # a withdrawn draft is frozen too
            conn.execute("UPDATE change_requests SET status = '申請済み' WHERE id = ?", (rid,))
        _event(conn, user, rid, "取り下げ", config.CR_WITHDRAWN, reason=reason)
    _ok(conn, user, "change_withdraw", rid, "取り下げた（記録は残る）", req["to_version_id"])
    return config.CR_WITHDRAWN


# ---- before the test ---------------------------------------------------------------------------------------------------

@_refusals_logged("change_pre_decide")
def pre_decide(conn: sqlite3.Connection, user: User, rid: int, approve: bool, reason: str = "") -> str:
    """Approve (the test then runs by itself) or reject before the test. An administrator who is not the requester."""
    op = "事前承認" if approve else "事前却下"
    roles.require(conn, user, "change_pre_decide", env=PROD)
    req = _req(conn, rid)
    _not_requester(conn, user, req, "change_pre_decide")
    _need_state(conn, req, (config.CR_WAIT_PRE,), op)
    if not approve and not reason.strip():
        raise ChangeError("却下には理由の文が要る")
    with conn:
        _event(conn, user, rid, op, config.CR_WAIT_TEST if approve else config.CR_REJECTED, reason=reason,
               version_id=req["to_version_id"])
    _ok(conn, user, "change_pre_decide", rid, op, req["to_version_id"])
    if not approve:
        return config.CR_REJECTED
    return _run_test(conn, user, rid)


def _run_test(conn: sqlite3.Connection, user: User, rid: int) -> str:
    """Run the test and record the result. A failure caused by the change is a result (it is recorded and waits for
    a decision); a fault of the tool is not (changecheck.ToolError goes on up and the request stays at 試験待ち)."""
    req = _req(conn, rid)
    try:
        out = changecheck.run(conn, user, req["to_version_id"], PROD)
    except changecheck.ToolError as e:
        _denied(conn, user, "change_test", rid, f"試験の道具が止まった（試験済みにしない）: {e}")
        raise
    with conn:
        cur = conn.execute(
            "INSERT INTO change_tests(request_id, run_at, run_by, base_version_id, to_version_id, base_sha256,"
            " to_sha256, prod_copy_at, samples_json, result_status, result_json) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (rid, now_iso(), user.id, out.result["base_version_id"], req["to_version_id"], out.base_sha256,
             out.to_sha256, out.prod_copy_at, json.dumps(out.samples, ensure_ascii=False, sort_keys=True), out.status,
             json.dumps(out.result, ensure_ascii=False)))
        tid = int(cur.lastrowid)
        _event(conn, user, rid, "試験", config.CR_TESTED, reason=f"試験の結果 {tid}: {out.status}",
               version_id=req["to_version_id"], env="使い捨ての DB")
    _ok(conn, user, "change_test", rid, f"試験の結果 {tid}: {out.status}", req["to_version_id"])
    return config.CR_TESTED


@_refusals_logged("change_retest")
def retest(conn: sqlite3.Connection, user: User, rid: int) -> str:
    """Run the test again (also after the tool stopped). The approvals of the earlier result no longer count."""
    roles.require(conn, user, "change_retest", env=PROD)
    req = _req(conn, rid)
    st = _need_state(conn, req, (config.CR_WAIT_TEST, config.CR_TESTED, config.CR_RESULT_OK), "試験をやり直")
    if st != config.CR_WAIT_TEST:
        with conn:
            _event(conn, user, rid, "試験のやり直し", config.CR_WAIT_TEST,
                   reason="前の結果への承認は数えない", version_id=req["to_version_id"])
    return _run_test(conn, user, rid)


# ---- the result ---------------------------------------------------------------------------------------------------------

def latest_test(conn: sqlite3.Connection, rid: int) -> sqlite3.Row | None:
    return conn.execute("SELECT * FROM change_tests WHERE request_id = ? ORDER BY id DESC LIMIT 1", (rid,)).fetchone()


def approvals_of(conn: sqlite3.Connection, test_id: int) -> list[sqlite3.Row]:
    return conn.execute("SELECT * FROM change_approvals WHERE test_id = ? ORDER BY id", (test_id,)).fetchall()


def enough_approvals(approvals: list[tuple[int, frozenset[str]]]) -> bool:
    """Two different people: one holding 承認者 and one holding 管理者 (one person with both roles is one person)."""
    for uid_a, roles_a in approvals:
        if "承認者" not in roles_a:
            continue
        for uid_b, roles_b in approvals:
            if uid_b != uid_a and "管理者" in roles_b:
                return True
    return False


def _approval_pairs(rows: list[sqlite3.Row]) -> list[tuple[int, frozenset[str]]]:
    return [(r["user_id"], frozenset(x for x in r["roles"].split(",") if x)) for r in rows]


@_refusals_logged("change_result_decide")
def decide_result(conn: sqlite3.Connection, user: User, rid: int, approve: bool, reason: str = "", *,
                  test_id: int | None = None) -> str:
    """Approve or reject the result of the latest test. Not the requester. Either of the two kinds of person may reject.
    `test_id` is the test result the person has on his screen (the screens always send it; 0 = not sent): when the test
    ran again after he opened the page, the decision is refused, so an approval never counts for a result he has not
    read."""
    op = "結果の承認" if approve else "結果の却下"
    roles.require(conn, user, "change_result_decide", env=PROD)
    req = _req(conn, rid)
    _not_requester(conn, user, req, "change_result_decide")
    _need_state(conn, req, (config.CR_TESTED,), op)
    test = latest_test(conn, rid)
    if test is None:
        raise ChangeError("試験の結果が無い")
    if test_id is not None and test_id != test["id"]:
        if not test_id:
            raise NotLatestResult("どの試験の結果に対する判断かが送られていない（画面を開き直す）")
        raise NotLatestResult(f"画面の試験の結果 {test_id} は最新ではない（最新は {test['id']}。試験が流し直された）。"
                              "画面を開き直して、最新の結果を読んでから決める")
    if not approve:
        if not reason.strip():
            raise ChangeError("却下には理由の文が要る")
        with conn:
            _event(conn, user, rid, op, config.CR_REJECTED, reason=reason, version_id=req["to_version_id"])
        _ok(conn, user, "change_result_decide", rid, f"{op}（試験の結果 {test['id']}）", req["to_version_id"])
        return config.CR_REJECTED
    if any(r["user_id"] == user.id for r in approvals_of(conn, test["id"])):
        raise ChangeError("この結果にはすでに承認している")
    with conn:
        conn.execute("INSERT INTO change_approvals(request_id, test_id, user_id, user_name, roles, approved_at)"
                     " VALUES (?, ?, ?, ?, ?, ?)",
                     (rid, test["id"], user.id, user.name, ",".join(sorted(user.roles)), now_iso()))
        done = enough_approvals(_approval_pairs(approvals_of(conn, test["id"])))
        _event(conn, user, rid, op, config.CR_RESULT_OK if done else config.CR_TESTED,
               reason=f"試験の結果 {test['id']} に承認" + ("。2 人そろった" if done else ""),
               version_id=req["to_version_id"])
    _ok(conn, user, "change_result_decide", rid,
        f"{op}（試験の結果 {test['id']}）" + ("。2 人そろった" if done else ""), req["to_version_id"])
    return config.CR_RESULT_OK if done else config.CR_TESTED


# ---- applying and confirming ---------------------------------------------------------------------------------------------

@_refusals_logged("change_apply")
def apply(conn: sqlite3.Connection, user: User, rid: int, version_id: int | None = None) -> str:
    """Point production at the version. `version_id` is the version the person sees on the screen (default: the
    version of the request). The move is one `UPDATE ... WHERE version_id IS <the version the test was run against>`
    in one transaction: if another request got in first, nothing moves and the request goes back to 試験待ち."""
    roles.require(conn, user, "change_apply", env=PROD)
    req = _req(conn, rid)
    _not_requester(conn, user, req, "change_apply")
    _need_state(conn, req, (config.CR_RESULT_OK,), "適用")
    test = latest_test(conn, rid)
    if test is None or not enough_approvals(_approval_pairs(approvals_of(conn, test["id"]))):
        msg = "最新の試験の結果に、承認者と管理者の 2 人の承認がそろっていない"
        _denied(conn, user, "change_apply", rid, msg)
        raise ChangeError(msg)
    target = version_id if version_id is not None else req["to_version_id"]
    row = conn.execute("SELECT normalized_json FROM config_versions WHERE id = ?", (target,)).fetchone()
    if row is None:
        raise ChangeError(f"設定の版 {target} は無い")
    import hashlib
    sha_now = hashlib.sha256(row["normalized_json"].encode("utf-8")).hexdigest()    # computed again, not read
    if sha_now != test["to_sha256"]:
        msg = f"試験した版と違う（版 {target} の sha256 が、試験の結果に記録した値と合わない）。適用しない"
        _denied(conn, user, "change_apply", rid, msg)
        raise ChangeError(msg)
    if target != req["to_version_id"]:
        msg = f"申請の版は {req['to_version_id']} で、{target} ではない。適用しない"
        _denied(conn, user, "change_apply", rid, msg)
        raise ChangeError(msg)
    check_not_reopening(conn, user, json.loads(row["normalized_json"]), "change_apply", rid)
    with conn:
        moved = versions.point_environment(conn, PROD, target, user.id, expect=test["base_version_id"],
                                           change_request_id=rid)
        if moved:
            _event(conn, user, rid, "適用", config.CR_APPLIED, version_id=target,
                   reason=f"本番の版: {_ver(test['base_version_id'])} -> {_ver(target)}")
    if not moved:
        with conn:
            _event(conn, user, rid, "適用の拒否", config.CR_WAIT_TEST,
                   reason="本番の版が、試験をした時から変わっていた。試験をやり直す", version_id=target)
        msg = "本番の版が試験の時から変わっている（先の申請が適用された）。何も変えず、試験のやり直しに戻した"
        _denied(conn, user, "change_apply", rid, msg)
        raise StaleTest(msg)
    _ok(conn, user, "change_apply", rid, f"本番を版 {target} にした（試験の時の本番: {_ver(test['base_version_id'])}）", target)
    return config.CR_APPLIED


@_refusals_logged("change_confirm")
def confirm(conn: sqlite3.Connection, user: User, rid: int, ok: bool, reason: str = "") -> str:
    """The requester says whether the change came out as asked. 確認できない does not close the request."""
    op = "確認" if ok else "確認できない"
    roles.require(conn, user, "change_confirm", env=PROD)
    req = _req(conn, rid)
    if req["requester_id"] != user.id:
        _denied(conn, user, "change_confirm", rid, "確認は申請者だけ")
        raise NotAllowedToChange("確認できるのは、その変更を申請した人だけ（求めた通りに出たかを見る人）")
    _need_state(conn, req, (config.CR_APPLIED, config.CR_CANNOT_CONFIRM), op)
    if not ok and not reason.strip():
        raise ChangeError("確認できない時は、その理由の文が要る")
    new = config.CR_CONFIRMED if ok else config.CR_CANNOT_CONFIRM
    with conn:
        _event(conn, user, rid, op, new, reason=reason, version_id=req["to_version_id"])
    _ok(conn, user, "change_confirm", rid, op + ("（閉じた）" if ok else "（閉じずに置く）"), req["to_version_id"])
    return new


@_refusals_logged("change_create")
def draft_rollback_of(conn: sqlite3.Connection, user: User, rid: int, justification: str) -> int:
    """From a request that cannot be confirmed: a draft that puts back the version of its back-out plan (item 15: the
    version the requester named, or - by default - the version that was in production when it was put forward)."""
    roles.require(conn, user, "change_create", env=PROD)          # the role first: whatever the state
    req = _req(conn, rid)
    _need_state(conn, req, (config.CR_CANNOT_CONFIRM,), "戻しの下書きを作")
    back = req["rollback_version_id"] if req["rollback_version_id"] is not None else req["from_version_id"]
    if back is None:
        raise ChangeError("この申請の前は「版なし」で、戻し先の版が無い")
    return create_rollback(conn, user, back, justification=justification, source_request_id=rid)


# ---- reading --------------------------------------------------------------------------------------------------------------

def list_requests(conn: sqlite3.Connection, user: User) -> list[dict[str, Any]]:
    roles.require(conn, user, "view", env=PROD)
    out = []
    for r in conn.execute(
            "SELECT c.*, u.name AS requester_name FROM change_requests c JOIN users u ON u.id = c.requester_id"
            " ORDER BY c.id DESC"):
        d = dict(r)
        d["state"] = state_of(conn, r["id"])
        out.append(d)
    return out


def _person_marks(conn: sqlite3.Connection, req: sqlite3.Row, names: list[str]) -> dict[str, dict[str, Any]]:
    out = {}
    for n in dict.fromkeys(names):
        u = conn.execute("SELECT created_at FROM users WHERE name = ?", (n,)).fetchone()
        created = u["created_at"] if u else ""
        out[n] = {"created_at": created, "after_request": bool(created) and created > req["created_at"]}
    return out


def detail(conn: sqlite3.Connection, user: User, rid: int) -> dict[str, Any]:
    roles.require(conn, user, "view", env=PROD)
    req = _req(conn, rid)
    d: dict[str, Any] = {"req": dict(req), "state": state_of(conn, rid)}
    d["requester_name"] = conn.execute("SELECT name FROM users WHERE id = ?", (req["requester_id"],)).fetchone()["name"]
    d["diff"] = json.loads(req["diff_json"]) if req["diff_json"] else None
    d["impact"] = json.loads(req["impact_json"]) if req["impact_json"] else None
    d["security"] = json.loads(req["security_json"]) if req["security_json"] else None
    d["events"] = [dict(e) for e in conn.execute(
        "SELECT * FROM change_request_events WHERE request_id = ? ORDER BY id", (rid,))]
    tests = [dict(t) for t in conn.execute("SELECT * FROM change_tests WHERE request_id = ? ORDER BY id", (rid,))]
    for t in tests:
        t["result"] = json.loads(t["result_json"])
        t["samples"] = json.loads(t["samples_json"])
    d["tests"] = tests
    d["test"] = tests[-1] if tests else None
    current = approvals_of(conn, d["test"]["id"]) if d["test"] else []
    d["approvals"] = [dict(a) for a in current]
    d["approver_ok"] = any("承認者" in a["roles"].split(",") and a["user_id"] != req["requester_id"] for a in current)
    d["enough"] = enough_approvals(_approval_pairs(current))
    d["marks"] = _person_marks(conn, req, [e["user_name"] for e in d["events"]] + [a["user_name"] for a in current])
    widened = (d["security"] or {}).get("widened", {})
    d["widened_among_approvers"] = {a["user_name"]: widened[a["user_name"]] for a in d["approvals"]
                                    if a["user_name"] in widened}
    d["versions"] = {}
    for key in ("from_version_id", "to_version_id", "rollback_version_id"):
        vid = req[key]
        if vid is not None:
            v = conn.execute("SELECT id, sha256, loaded_at, loaded_by FROM config_versions WHERE id = ?", (vid,)).fetchone()
            d["versions"][key] = dict(v) if v else None
    d["production_version_id"] = production_version(conn)
    return d


def allowed_actions(user: User, d: dict[str, Any]) -> list[str]:
    """What the screen offers this person at this state (the server still checks every one when it is posted)."""
    req, st = d["req"], d["state"]
    me_req = req["requester_id"] == user.id
    admin, appr = "管理者" in user.roles, "承認者" in user.roles
    acts: list[str] = []
    if st == config.CR_DRAFT and admin and me_req:
        acts.append("submit")
    if st in (config.CR_DRAFT, config.CR_WAIT_PRE, config.CR_WAIT_TEST, config.CR_TESTED, config.CR_RESULT_OK) \
            and admin and me_req:
        acts.append("withdraw")
    if st == config.CR_WAIT_PRE and admin and not me_req:
        acts += ["pre_approve", "pre_reject"]
    if st in (config.CR_WAIT_TEST, config.CR_TESTED, config.CR_RESULT_OK) and admin:
        acts.append("retest")
    if st == config.CR_TESTED and (admin or appr) and not me_req:
        if not any(a["user_id"] == user.id for a in d["approvals"]):
            acts.append("result_approve")
        acts.append("result_reject")
    if st == config.CR_RESULT_OK and admin and not me_req:
        acts.append("apply")
    if st in (config.CR_APPLIED, config.CR_CANNOT_CONFIRM) and admin and me_req:
        acts += ["confirm", "cannot_confirm"] if st == config.CR_APPLIED else ["confirm"]
    if st == config.CR_CANNOT_CONFIRM and admin:
        acts.append("rollback")
    return acts


CSV_HEADER = ["番号", "作成日", "題名", "種類", "申請者", "状態", "元の版（申請の時の本番）", "上げる版", "理由",
              "急ぎの度合い", "期待する機能の影響", "しない時の影響", "戻し先の版", "変わる行の数", "統制に関わる変更"]


def to_csv(conn: sqlite3.Connection, user: User) -> str:
    roles.require(conn, user, "audit_csv", env=PROD)
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n", quoting=csv.QUOTE_ALL)
    w.writerow(CSV_HEADER)
    cell = audit._cell
    for r in reversed(list_requests(conn, user)):
        diff = json.loads(r["diff_json"]) if r["diff_json"] else []
        sec = json.loads(r["security_json"]) if r["security_json"] else {}
        w.writerow([r["id"], r["created_at"], cell(r["title"]), r["kind"], cell(r["requester_name"]), r["state"],
                    r["from_version_id"] if r["from_version_id"] is not None else "版なし", r["to_version_id"],
                    cell(r["justification"]), r["urgency"], cell(r["functional_impact"]), cell(r["impact_if_not"]),
                    r["rollback_version_id"] if r["rollback_version_id"] is not None else "版なし", len(diff),
                    "あり" if sec.get("any") else ("なし" if sec else "")])
    return out.getvalue()
