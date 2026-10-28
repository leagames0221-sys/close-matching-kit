"""Foundation of the second release: the numbered migration, held rows, balances and their continuity, the
period-end balance, the schema 2 workbook."""
import json
import sqlite3
from datetime import date as D

import pytest

from close_kit import balance, config, db, ingest, journals, matchrun, reasons, versions, workbook
from conftest import load_workbook_into
from tools import make_samples, sample_builders as sb
from close_helpers import MS, bank_interest_row, group_with, txn

# ---- migration ---------------------------------------------------------------------------------------------


def old_database(tmp_path):
    """A database exactly as the first release left it (migration 1 only), with rows in the tables to be rebuilt."""
    c = sqlite3.connect(str(tmp_path / "old.sqlite3"))
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys = ON")
    c.executescript(db.SCHEMA)
    c.execute("INSERT INTO users(name, pw_hash, roles, created_at) VALUES ('u', 'x', '管理者', 'now')")
    c.execute("INSERT INTO config_versions(schema_version, normalized_json, sha256, source_bytes, loaded_by, loaded_at)"
              " VALUES ('1', '{}', 'h', x'00', 1, 'now')")
    c.execute("INSERT INTO runs(kind, env, version_id, created_by, created_at) VALUES ('match', 'sandbox', 1, 1, 'now')")
    c.execute("INSERT INTO batches(run_id, env, version_id, source, account, fmt, sha256, size, row_count)"
              " VALUES (1, 'sandbox', 1, 's', 'a', 'csv', 'h', 1, 2)")
    for i in (1, 2):
        c.execute("INSERT INTO transactions(batch_id, run_id, env, version_id, source, account, booking_date,"
                  " amount_minor, src_row, ext_id) VALUES (1, 1, 'sandbox', 1, 's', 'a', '2026-10-01', ?, ?, ?)",
                  (i * 100, i, str(i)))
    c.execute("INSERT INTO match_groups(id, run_id, env, version_id, match_set, kind, state, created_at)"
              " VALUES (7, 1, 'sandbox', 1, 'm', '自動', '完了', 'now')")
    c.execute("INSERT INTO match_group_rows(group_id, txn_id, side) VALUES (7, 1, 'L'), (7, 2, 'R')")
    c.execute("INSERT INTO match_rows(env, match_set, txn_id, side, klass, group_id, run_id)"
              " VALUES ('sandbox', 'm', 1, 'L', '自動', 7, 1), ('sandbox', 'm', 2, 'R', '自動', 7, 1)")
    c.commit()
    return c


def schema_dump(c):
    return [tuple(r) for r in c.execute("SELECT type, name, sql FROM sqlite_master ORDER BY type, name")]


def test_migration_keeps_every_row_and_id_and_running_it_again_changes_nothing(tmp_path):
    c = old_database(tmp_path)
    assert "仕訳の承認待ち" not in c.execute("SELECT sql FROM sqlite_master WHERE name = 'match_groups'").fetchone()[0]
    db.migrate(c)
    assert [r[0] for r in c.execute("SELECT version FROM schema_migrations ORDER BY version")] == [1, 2, 3, 4]
    assert [tuple(r) for r in c.execute("SELECT id, state FROM match_groups")] == [(7, "完了")]
    assert [tuple(r) for r in c.execute("SELECT txn_id, klass, group_id FROM match_rows ORDER BY txn_id")] == \
        [(1, "自動", 7), (2, "自動", 7)]
    assert c.execute("SELECT COUNT(*) FROM match_group_rows").fetchone()[0] == 2
    assert c.execute("PRAGMA foreign_keys").fetchone()[0] == 1
    assert c.execute("PRAGMA foreign_key_check").fetchall() == []
    before = schema_dump(c)
    rows_before = [tuple(r) for r in c.execute("SELECT * FROM match_rows ORDER BY txn_id")]
    db.migrate(c)
    db.migrate(c)
    assert schema_dump(c) == before
    assert [tuple(r) for r in c.execute("SELECT * FROM match_rows ORDER BY txn_id")] == rows_before
    assert c.execute("SELECT COUNT(*) FROM schema_migrations").fetchone()[0] == 4


def test_migration_4_adds_the_value_for_a_journal_the_tool_drafted_and_keeps_rows_triggers_and_links(tmp_path, monkeypatch):
    c = old_database(tmp_path)
    full = list(db.MIGRATIONS)
    monkeypatch.setattr(db, "MIGRATIONS", full[:2])                       # a database as the third release left it
    db.migrate(c)
    assert [r[0] for r in c.execute("SELECT version FROM schema_migrations ORDER BY version")] == [1, 2, 3]
    c.execute("INSERT INTO journals(id, env, number, entry_date, period, memo, method, status, debit_total, created_by,"
              " created_at, version_id) VALUES (5, 'sandbox', 'JE-000005', '2026-10-31', '2026-10', 'm', '手入力',"
              " '承認待ち', 700, 1, 'now', 1)")
    c.execute("INSERT INTO journal_lines(journal_id, line_no, account_code, debit, credit) VALUES (5, 1, '6100', 700, 0),"
              " (5, 2, '1100', 0, 700)")
    c.execute("UPDATE journals SET status = '承認済み' WHERE id = 5")
    c.execute("INSERT INTO posting_log(env, journal_id, action, file_id, sha256, by_user, at, note)"
              " VALUES ('sandbox', 5, '出力', NULL, 'h', 1, 'now', '')")
    c.commit()
    with pytest.raises(sqlite3.IntegrityError):                            # before: the new value is refused
        c.execute("INSERT INTO journals(env, number, entry_date, period, method, status, debit_total, created_by,"
                  " created_at, version_id) VALUES ('sandbox', 'JE-X', '2026-10-31', '2026-10', '照合の差額', '下書き', 1,"
                  " 1, 'now', 1)")
    monkeypatch.setattr(db, "MIGRATIONS", full)
    db.migrate(c)
    assert [r[0] for r in c.execute("SELECT version FROM schema_migrations ORDER BY version")] == [1, 2, 3, 4]
    assert [tuple(r) for r in c.execute("SELECT id, number, status, method FROM journals")] ==         [(5, "JE-000005", "承認済み", "手入力")]
    assert c.execute("SELECT COUNT(*) FROM journal_lines WHERE journal_id = 5").fetchone()[0] == 2
    assert c.execute("SELECT COUNT(*) FROM posting_log WHERE journal_id = 5").fetchone()[0] == 1
    assert c.execute("PRAGMA foreign_keys").fetchone()[0] == 1
    assert c.execute("PRAGMA foreign_key_check").fetchall() == []
    for sql in ("UPDATE journals SET memo = 'x' WHERE id = 5", "DELETE FROM journals WHERE id = 5",
                "UPDATE journal_lines SET debit = 1 WHERE journal_id = 5"):   # the triggers were made again
        with pytest.raises(sqlite3.IntegrityError):
            c.execute(sql)
    c.execute("INSERT INTO journals(env, number, entry_date, period, method, status, debit_total, created_by, created_at,"
              " version_id) VALUES ('sandbox', 'JE-X', '2026-10-31', '2026-10', '照合の差額', '下書き', 1, 1, 'now', 1)")
    with pytest.raises(sqlite3.IntegrityError):                            # still a CHECK list
        c.execute("INSERT INTO journals(env, number, entry_date, period, method, status, debit_total, created_by,"
                  " created_at, version_id) VALUES ('sandbox', 'JE-Y', '2026-10-31', '2026-10', '何か', '下書き', 1, 1,"
                  " 'now', 1)")
    c.rollback()
    before = schema_dump(c)
    db.migrate(c)
    assert schema_dump(c) == before


def test_migration_widens_the_check_lists_and_keeps_them_checks(tmp_path):
    c = old_database(tmp_path)
    with pytest.raises(sqlite3.IntegrityError):  # before: the new state is refused
        c.execute("UPDATE match_groups SET state = '仕訳の承認待ち' WHERE id = 7")
    db.migrate(c)
    c.execute("UPDATE match_groups SET state = '仕訳の承認待ち' WHERE id = 7")
    c.execute("UPDATE match_rows SET klass = '処理済み', group_id = NULL WHERE txn_id = 1")
    c.execute("UPDATE match_rows SET klass = '転記済み' WHERE txn_id = 2")
    with pytest.raises(sqlite3.IntegrityError):
        c.execute("UPDATE match_groups SET state = '何か' WHERE id = 7")
    with pytest.raises(sqlite3.IntegrityError):
        c.execute("UPDATE match_rows SET klass = '何か' WHERE txn_id = 1")
    cols = {r[1] for r in c.execute("PRAGMA table_info(batches)")}
    assert {"opening_minor", "closing_minor"} <= cols
    assert "balance_minor" in {r[1] for r in c.execute("PRAGMA table_info(transactions)")}


def test_every_new_table_carries_the_environment(conn):
    for t in ("journals", "adjusting_items", "posting_files", "posting_log", "posted_links", "task_instances",
              "task_events"):
        assert "env" in {r[1] for r in conn.execute(f"PRAGMA table_info({t})")}, t


# ---- workbook schema 2 ---------------------------------------------------------------------------------------

def parse(sheets=None, **kw):
    return workbook.parse_workbook(sb.build_workbook(sheets, **kw))


def errs(res):
    return [str(e) for e in res.errors]


def close_sheets(**kw):
    return make_samples.close_sheets(**kw)


def test_the_closing_sample_workbook_reads_and_has_the_new_content():
    res = workbook.parse_workbook(make_samples.close_workbook())
    assert res.ok, errs(res)
    cfg = res.config
    assert cfg["description"]["company_code"] == "1000" and cfg["description"]["journal_type"] == "SA"
    ms = cfg["match_sets"][0]
    assert (ms["bank_side"], ms["deposit_account"], ms["reconcile"]) == ("右", "1100", True)
    assert [r["lower"] for r in cfg["approval_routes"]] == [0, 100000]
    assert [t["id"] for t in cfg["tasks"]] == ["T01", "T02", "T03", "T04", "T05", "T06", "T07"]
    assert {t["counting"] for t in cfg["tasks"]} == {"営業日", "暦日"}
    assert cfg["company_holidays"] == [{"date": "2026-11-06", "name": "創立記念日（架空）"}]


def test_a_version_1_workbook_is_refused_as_too_old():
    sheets = sb.default_sheets()
    sheets["説明"] = [["1", "社", "", D(2026, 10, 24), "1000", "SA"]]
    res = parse(sheets)
    assert not res.ok and any("schema の版が古い" in m for m in errs(res))
    # a real version 1 workbook (no new sheets, no new columns) says the same
    from openpyxl import Workbook
    import io
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("説明")
    ws.append(["schemaの版", "会社名", "作成者", "更新日"])
    ws.append(["1", "社", "", D(2026, 10, 24)])
    out = io.BytesIO()
    wb.save(out)
    res = workbook.parse_workbook(out.getvalue())
    assert not res.ok and any("schema の版が古い" in m for m in errs(res))


@pytest.mark.parametrize("edit, needle", [
    (lambda s: s["勘定科目"].append(["12345678901", "長すぎる", "資産"]), "10 文字までにする"),
    (lambda s: s.__setitem__("説明", [["2", "社", "", D(2026, 10, 24), "10000", "SA"]]), "4 文字までにする"),
    (lambda s: s.__setitem__("説明", [["2", "社", "", D(2026, 10, 24), "1000", "SAP"]]), "2 文字までにする"),
    (lambda s: s["照合セット"][0].__setitem__(5, "4100"), "区分が資産の勘定にする"),
    (lambda s: s["照合セット"][0].__setitem__(5, "9999"), "勘定科目に無いコード"),
    (lambda s: s["照合セット"][0].__setitem__(4, "上"), "左 / 右"),
    # written as they are into the posting file, so a formula-looking code is refused at the entrance
    (lambda s: s["勘定科目"].append(["-1", "数式", "費用"]), "数式になる文字で始めない"),   # "=" is a formula cell already
    (lambda s: s["勘定科目"].append(["＋9", "全角の数式", "費用"]), "数式になる文字で始めない"),
    (lambda s: s.__setitem__("説明", [["2", "社", "", D(2026, 10, 24), "@1", "SA"]]), "数式になる文字で始めない"),
    (lambda s: s.__setitem__("説明", [["2", "社", "", D(2026, 10, 24), "1000", "-A"]]), "数式になる文字で始めない"),
    # a journal from the deposit account to itself would leave the difference where it is
    (lambda s: s["理由コード"][0].__setitem__(3, "1100"), "相手勘定にできない"),
])
def test_match_set_and_account_columns_are_checked(edit, needle):
    s = close_sheets()
    edit(s)
    res = parse(s)
    assert not res.ok and any(needle in m for m in errs(res)), errs(res)


def test_a_balance_match_set_needs_balances_and_one_unit():
    s = close_sheets()
    s["データ源"][0][11] = ""  # the ledger CSV has no balance column any more
    assert any("残高が要る" in m for m in errs(parse(s)))
    s = close_sheets()
    s["データ源"][0][4] = 100
    assert any("最小単位の倍率が違う" in m for m in errs(parse(s)))
    s = close_sheets()
    s["照合セット"][0][6] = "いいえ"      # not used for the balance: no balance needed, a mismatch is fine
    s["作業一覧"] = []
    s["データ源"][0][11] = ""
    s["データ源"][0][4] = 100
    assert parse(s).ok


def test_deposit_accounts_of_balance_match_sets_do_not_overlap():
    s = close_sheets()
    s["データ源"] += [["帳簿2", "csv", "", "帳簿の借方を + にする", 1, "西暦", "取引ID", "日付", "金額", "参照", "摘要", "残高"],
                    ["銀行2", "zengin", "7654321", "入金を + にする", 1, "令和", "", "", "", "", "", ""]]
    s["照合セット"].append(["預金照合2", "帳簿2", "銀行2", "はい", "右", "1100", "はい"])
    assert any("預金の勘定が重なる" in m for m in errs(parse(s)))
    s["勘定科目"].append(["1110", "定期預金", "資産"])
    s["照合セット"][1][5] = "1110"
    assert parse(s).ok, errs(parse(s))


@pytest.mark.parametrize("routes, needle", [
    ([[0, 100, "a"], [200, "", "b"]], "隙間がある"),
    ([[0, 100, "a"], [50, "", "b"]], "重なっている"),
    ([[0, "", "a"], [100, "", "b"]], "上限が空なのは最後の行だけ"),
    ([[0, 100, "a"], [100, 200, "b"]], "最後の行の上限は空"),
    ([[10, "", "a"]], "最初の幅の下限は 0"),
    ([[0, "", "、"]], "名前が 1 人も無い"),
    ([], "行が 1 つも無い"),
])
def test_approval_route_ranges_must_have_no_gap_no_overlap_and_an_open_last_row(routes, needle):
    s = close_sheets()
    s["承認ルート"] = routes
    assert any(needle in m for m in errs(parse(s))), errs(parse(s))


def test_the_range_of_days_from_the_period_end_is_a_setting_of_the_workbook(conn, users):
    """Oracle Close Manager: "You can specify the number of days a template tracks before and after the Zero day." The
    two columns of the 説明 sheet set the range; empty means -20 to 40."""
    def with_range(lo, hi, n):
        s = close_sheets()
        s["説明"] = [[*s["説明"][0][:6], lo, hi]]
        s["作業一覧"] = [["A", "名", "繰り返し", "", "暦日", n, "", "a", "b", "なし", "", "はい"]]
        return parse(s)
    ok = with_range(-60, 90, 90)
    assert ok.ok, errs(ok)
    assert (ok.config["description"]["task_offset_min"], ok.config["description"]["task_offset_max"]) == (-60, 90)
    assert any("-60 から 90 の整数にする" in m for m in errs(with_range(-60, 90, 91)))
    assert any("-60 から 90 の整数にする" in m for m in errs(with_range(-60, 90, -61)))
    assert with_range(-60, 90, -60).ok
    assert any("-20 から 40" in m for m in errs(with_range("", "", 41)))           # empty = the default range
    assert with_range("", "", 40).ok and with_range(-5, "", 40).ok
    assert any("0 以下にする" in m for m in errs(with_range(1, 40, 1)))            # day 0 is always inside
    assert any("0 以上にする" in m for m in errs(with_range(-20, -1, -5)))
    # a due date far outside the default range is made from the setting
    load_workbook_into(conn, users.admin, sb.build_workbook(
        {**{k: v for k, v in close_sheets().items()},
         "説明": [[*close_sheets()["説明"][0][:6], -20, 90]],
         "作業一覧": [["A", "長い作業", "繰り返し", "", "暦日", 75, "", "demo-maker", "demo-approver", "なし", "", "はい"]]}))
    from close_kit import tasks
    tasks.generate(conn, users.admin, "2026-10")
    assert conn.execute("SELECT due_date FROM task_instances WHERE task_id = 'A'").fetchone()[0] == "2027-01-14"


def test_task_sheet_rules():
    base = close_sheets()["作業一覧"]

    def with_tasks(rows):
        s = close_sheets()
        s["作業一覧"] = rows
        return errs(parse(s))
    t = lambda i, **kw: [i, "名", kw.get("type", "繰り返し"), kw.get("period", ""), "営業日", kw.get("n", 1), kw.get("pre", ""),  # noqa: E731
                         kw.get("prep", "a"), kw.get("rev", "b"), kw.get("cond", "なし"), kw.get("ms", ""), "はい"]
    assert not with_tasks([t("A"), t("B", pre="A")])
    assert any("輪になっている" in m for m in with_tasks([t("A", pre="B"), t("B", pre="A")]))
    assert any("自分自身" in m for m in with_tasks([t("A", pre="A")]))
    assert any("作業 ID に無い名前" in m for m in with_tasks([t("A", pre="Z")]))
    assert any("作成者と確認者が同じ人" in m for m in with_tasks([t("A", prep="a", rev="a")]))
    assert any("対象の期間が必要" in m for m in with_tasks([t("A", type="プロジェクト")]))
    assert any("会計期間に無い名前" in m for m in with_tasks([t("A", type="臨時", period="2030-01")]))
    assert not with_tasks([t("A", type="プロジェクト", period="2026-10")])
    assert any("繰り返しの作業では対象の期間は使わない" in m for m in with_tasks([t("A", period="2026-10")]))
    assert any("-20 から 40" in m for m in with_tasks([t("A", n=41)]))
    assert any("-20 から 40" in m for m in with_tasks([t("A", n=-21)]))
    assert not with_tasks([t("A", n=40), t("B", n=-20)])
    assert any("照合セットが必要" in m for m in with_tasks([t("A", cond="未処理 0 件")]))
    assert any("照合セットに無い名前" in m for m in with_tasks([t("A", cond="未処理 0 件", ms="無い")]))
    assert not with_tasks([t("A", cond="残高一致の差 0", ms="預金照合")])
    assert any("作業 ID が重複" in m for m in with_tasks([t("A"), t("A")]))
    assert base  # the sample itself is valid


def test_the_condition_balance_difference_needs_a_balance_match_set():
    s = close_sheets()
    s["照合セット"][0][6] = "いいえ"
    s["データ源"][0][11] = ""
    assert any("「残高一致に使う」が「はい」ではない" in m for m in errs(parse(s)))  # task T02 asks for it


def test_a_version_that_reopens_a_closed_period_is_refused(conn, users):
    load_workbook_into(conn, users.admin, make_samples.close_workbook())
    s = close_sheets()
    s["会計期間"][0][3] = "いいえ"   # 2026-09 was closed
    data = sb.build_workbook(s)
    res = workbook.parse_workbook(data)
    assert res.ok
    with pytest.raises(versions.VersionRejected) as e:
        versions.save_version(conn, users.admin, res, data)
    assert "2026-09" in str(e.value)
    assert conn.execute("SELECT COUNT(*) FROM config_versions").fetchone()[0] == 1
    log = conn.execute("SELECT outcome, detail FROM audit_log WHERE action = 'workbook_load' ORDER BY id DESC"
                       ).fetchone()
    assert log["outcome"] == "拒否" and "2026-09" in log["detail"]
    # closing a further period is fine
    s = close_sheets()
    s["会計期間"][1][3] = "はい"
    load_workbook_into(conn, users.admin, sb.build_workbook(s))


def test_a_stored_version_of_schema_1_stops_the_new_features(conn, users):
    old = {"description": {"schema_version": "1"}, "periods": [], "accounts": [], "sources": [], "match_sets": [],
           "reason_codes": [], "pass_rules": []}
    with conn:
        conn.execute("INSERT INTO config_versions(schema_version, normalized_json, sha256, source_bytes, loaded_by,"
                     " loaded_at) VALUES ('1', ?, 'h', x'00', ?, 'now')", (json.dumps(old), users.admin.id))
        conn.execute("UPDATE environments SET version_id = 1 WHERE name = 'sandbox'")
    with pytest.raises(versions.SchemaTooOld) as e:
        journals.create_manual(conn, users.maker, "2026-10-01", "", [])
    assert "ワークブックを版 2 で読み直す" in str(e.value)
    with pytest.raises(versions.SchemaTooOld):
        reasons.apply_reason(conn, users.maker, "x", "R01", group_id=1)


# ---- balances, continuity, the period-end balance -------------------------------------------------------------

def test_balances_are_stored_on_bundles_and_rows(conn, oct_world):
    b = {r["source"]: r for r in conn.execute("SELECT * FROM batches")}
    assert (b["銀行"]["opening_minor"], b["銀行"]["closing_minor"]) == (1_000_000, make_samples.CLOSE_BANK_OCT_CLOSING)
    assert (b["帳簿"]["opening_minor"], b["帳簿"]["closing_minor"]) == (1_000_000, make_samples.CLOSE_LEDGER_OCT_CLOSING)
    bal = [r["balance_minor"] for r in conn.execute("SELECT balance_minor FROM transactions WHERE source = '帳簿' ORDER BY id")]
    assert bal == [1330000, 1480000, 1475000, 1434000, 1504000]
    assert conn.execute("SELECT COUNT(*) FROM transactions WHERE source = '銀行' AND balance_minor IS NOT NULL"
                        ).fetchone()[0] == 0   # zengin has balances on the bundle only


def test_the_next_bundle_must_start_where_the_last_one_ended(conn, oct_world):
    admin = oct_world.admin
    wrong = sb.build_zengin(make_samples.CLOSE_BANK_NOV, d_from=D(2026, 11, 1), d_to=D(2026, 11, 30), opening=999)
    r = ingest.ingest(conn, admin, "銀行", "wrong.txt", wrong)
    assert not r.ok
    assert any("999" in x and f"{make_samples.CLOSE_BANK_OCT_CLOSING:,}" in x and "直前の束" in x for x in r.reasons), r.reasons
    assert conn.execute("SELECT COUNT(*) FROM batches WHERE source = '銀行'").fetchone()[0] == 1
    assert ingest.ingest(conn, admin, "銀行", "ok.txt", make_samples.close_bank_nov()).ok


def test_a_bundle_filling_a_gap_must_fit_both_neighbours(conn, users):
    load_workbook_into(conn, users.admin, make_samples.close_workbook())
    admin = users.admin
    sep = sb.build_zengin([(D(2026, 9, 5), "in", 100, "ｱ", "", "S1")], d_from=D(2026, 9, 1), d_to=D(2026, 9, 30), opening=500)
    nov = make_samples.close_bank_nov()   # opens on 1,474,987
    assert ingest.ingest(conn, admin, "銀行", "nov.txt", nov).ok
    # September closes at 600 but November opens at 1,474,987: the September file is refused (the next bundle)
    r = ingest.ingest(conn, admin, "銀行", "sep.txt", sep)
    assert not r.ok and any("次の束" in x and "600" in x for x in r.reasons), r.reasons


def test_a_data_source_that_is_not_used_for_the_balance_is_not_chained(conn, users):
    s = close_sheets()
    s["照合セット"][0][6] = "いいえ"
    s["作業一覧"] = []
    load_workbook_into(conn, users.admin, sb.build_workbook(s))
    a = sb.build_zengin([(D(2026, 10, 2), "in", 1, "ｱ", "", "X")], d_from=D(2026, 10, 1), d_to=D(2026, 10, 31), opening=10)
    b = sb.build_zengin([(D(2026, 11, 2), "in", 1, "ｱ", "", "Y")], d_from=D(2026, 11, 1), d_to=D(2026, 11, 30), opening=999)
    assert ingest.ingest(conn, users.admin, "銀行", "a.txt", a).ok
    assert ingest.ingest(conn, users.admin, "銀行", "b.txt", b).ok


def csv_bundle(rows, opening=1000):
    """A ledger CSV with a balance column: rows = [(id, date, amount, balance_override)]."""
    out, bal = [], opening
    for i, d, amt, override in rows:
        bal += amt
        out.append([i, d, amt, i, "m", bal if override is None else override])
    return sb.build_csv(make_samples.CLOSE_LEDGER_HEADER, out)


def test_period_end_balance_from_a_bundle_that_crosses_the_period_end(conn, users):
    load_workbook_into(conn, users.admin, make_samples.close_workbook())
    bank = sb.build_zengin(
        [(D(2026, 10, 20), "in", 100, "ｱ", "", "A"), (D(2026, 10, 31), "out", 30, "ｲ", "", "B"),
         (D(2026, 11, 2), "in", 5000, "ｳ", "", "C")], d_from=D(2026, 10, 15), d_to=D(2026, 11, 10), opening=1000)
    assert ingest.ingest(conn, users.admin, "銀行", "b.txt", bank).ok
    # the bundle holds 10-31: opening + rows dated on or before 10-31 only (the row of 11-02 did not exist yet)
    assert balance.period_end_balance(conn, "sandbox", "銀行", "2026-10-01", "2026-10-31") == (1070, "")
    # at 11-30 the same bundle is the last one before it: opening + all its rows
    assert balance.period_end_balance(conn, "sandbox", "銀行", "2026-11-01", "2026-11-30") == (6070, "")
    # a period with no bundle that touches it: unknown
    amount, why = balance.period_end_balance(conn, "sandbox", "銀行", "2026-12-01", "2026-12-31")
    assert amount is None and "残高が分からない" in why


def test_csv_balance_column_and_the_calculation_must_agree(conn, users):
    load_workbook_into(conn, users.admin, make_samples.close_workbook())
    ok = csv_bundle([("T1", "2026-10-05", 100, None), ("T2", "2026-10-31", -30, None)])
    assert ingest.ingest(conn, users.admin, "帳簿", "l.csv", ok).ok
    assert balance.period_end_balance(conn, "sandbox", "帳簿", "2026-10-01", "2026-10-31") == (1070, "")


def test_csv_rows_out_of_date_order_make_the_two_ways_disagree(conn, users):
    load_workbook_into(conn, users.admin, make_samples.close_workbook())
    # the file order is T1 (10-31), T2 (11-02), T3 (10-30): the last row dated <= 10-31 in file order is T3, whose
    # printed balance already contains T2; the calculation (opening + rows <= 10-31) does not
    data = csv_bundle([("T1", "2026-10-31", 100, None), ("T2", "2026-11-02", 200, None), ("T3", "2026-10-30", 50, None)])
    assert ingest.ingest(conn, users.admin, "帳簿", "l.csv", data).ok
    amount, why = balance.period_end_balance(conn, "sandbox", "帳簿", "2026-10-01", "2026-10-31")
    assert amount is None and "CSV の行の残高" in why and "合わない" in why


def test_csv_balance_mismatch_inside_the_file_is_refused(conn, users):
    load_workbook_into(conn, users.admin, make_samples.close_workbook())
    bad = csv_bundle([("T1", "2026-10-05", 100, None), ("T2", "2026-10-06", -30, 5)])
    r = ingest.ingest(conn, users.admin, "帳簿", "l.csv", bad)
    assert not r.ok and any("残高が合わない" in x for x in r.reasons)


# ---- fixed states, held rows --------------------------------------------------------------------------------

def test_journal_pending_groups_are_fixed_in_a_rerun(conn, oct_world):
    u = oct_world
    gid = group_with(conn, "INV-2610-003")
    rows_before = conn.execute("SELECT txn_id, side FROM match_group_rows WHERE group_id = ? ORDER BY txn_id", (gid,)).fetchall()
    reasons.apply_reason(conn, u.maker, MS, "R01", group_id=gid)
    assert conn.execute("SELECT state FROM match_groups WHERE id = ?", (gid,)).fetchone()[0] == "仕訳の承認待ち"
    assert "仕訳の承認待ち" in matchrun.FIXED_STATES
    ids_before = [r["id"] for r in conn.execute("SELECT id FROM match_groups ORDER BY id")]
    matchrun.run_match(conn, u.admin, MS)
    assert [r["id"] for r in conn.execute("SELECT id FROM match_groups ORDER BY id")] == ids_before
    assert conn.execute("SELECT state FROM match_groups WHERE id = ?", (gid,)).fetchone()[0] == "仕訳の承認待ち"
    assert conn.execute("SELECT txn_id, side FROM match_group_rows WHERE group_id = ? ORDER BY txn_id", (gid,)).fetchall() == rows_before


def test_rows_with_a_journal_leave_the_remainder_and_the_manual_match(conn, oct_world):
    from close_kit import manual
    u = oct_world
    t = bank_interest_row(conn)
    res = reasons.apply_reason(conn, u.maker, MS, "R03", txn_id=t)
    other = txn(conn, "INV-2610-009", "帳簿")
    with pytest.raises(manual.ManualError) as e:
        manual.create_manual(conn, u.maker, MS, [other], [t])
    assert "仕訳" in str(e.value)
    s = matchrun.run_match(conn, u.admin, MS)
    row = conn.execute("SELECT klass FROM match_rows WHERE txn_id = ?", (t,)).fetchone()
    assert row is not None and row["klass"] == "未照合"                       # the row was not thrown away by the rerun
    assert s.fixed_rows >= 1
    journals.submit(conn, u.maker, res.id)
    journals.decide(conn, u.approver, res.id, True)
    assert conn.execute("SELECT klass FROM match_rows WHERE txn_id = ?", (t,)).fetchone()[0] == "処理済み"
    matchrun.run_match(conn, u.admin, MS)
    assert conn.execute("SELECT klass FROM match_rows WHERE txn_id = ?", (t,)).fetchone()[0] == "処理済み"
    assert conn.execute("SELECT COUNT(*) FROM match_groups WHERE id IN (SELECT group_id FROM match_rows WHERE txn_id = ?)", (t,)).fetchone()[0] == 0


def test_a_row_with_an_adjusting_item_stays_in_the_remainder_and_in_the_manual_match(conn, oct_world):
    from close_kit import manual
    u = oct_world
    t = txn(conn, "INV-2610-009", "帳簿")
    item = reasons.apply_reason(conn, u.maker, MS, "R04", txn_id=t)
    reasons.decide_item(conn, u.approver, item.id, True)
    matchrun.run_match(conn, u.admin, MS)
    assert conn.execute("SELECT klass FROM match_rows WHERE txn_id = ?", (t,)).fetchone()[0] == "未照合"
    partner = bank_interest_row(conn)
    gid = manual.create_manual(conn, u.maker, MS, [t], [partner])           # not refused: it is still in the remainder
    assert gid
